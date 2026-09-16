from io import BytesIO
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.core.config import get_settings
from backend.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as value:
        yield value


def _headers(client, username):
    response = client.post("/auth/login", data={
        "username": username,
        "password": get_settings().synthetic_user_password.get_secret_value(),
    })
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


@pytest.mark.integration
def test_mumbai_hot_leads_and_excel_export(client):
    headers = _headers(client, "bankuser0001")
    question = "What are the hot personal loan leads for Mumbai branch?"
    response = client.post("/ml/leads/query", headers=headers, json={"question": question})
    assert response.status_code == 200
    result = response.json()
    assert result["total_count"] == 12_000
    assert result["organization"] == "Mumbai Branch 01"
    assert result["source_table"] == "ml_campaign_leads"
    assert len(result["records"]) == 100
    assert all(row["lead_status"] == "HOT" for row in result["records"])
    assert {row["propensity_band"] for row in result["records"]} == {"HIGH", "VERY_HIGH"}
    first = result["records"][0]
    assert first["business_unit"] == "TIER_1"
    assert re.fullmatch(r"CUST\d{9}", first["customer_ref"])
    assert not first["customer_name"].startswith("Synthetic Customer")
    assert "average_ticket_size" not in first
    assert 10 <= first["interest_rate_pct"] <= 16
    assert any(row["business_unit"] == "PLTB" for row in result["records"])
    assert any(row["business_unit"] == "SALPL" for row in result["records"])

    export = client.post("/ml/leads/export", headers=headers, json={"question": question})
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content), read_only=True)
    assert sum(1 for _ in workbook.active.iter_rows(values_only=True)) == 12_001
    headers = next(workbook.active.iter_rows(values_only=True))
    assert {"business_unit", "interest_rate_pct"}.issubset(headers)
    assert "average_ticket_size" not in headers


@pytest.mark.integration
@pytest.mark.parametrize("question,source_table", [
    ("Show pre-approved personal loan customers in my branch", "ml_pl_underwriting_leads"),
    ("Who are the risky customers in my branch?", "ml_pl_risk_review_leads"),
    ("Show good customers tagged as bad by risk team in my branch", "ml_pl_risk_mismatch_leads"),
    ("Show high propensity PL customers in my branch", "ml_pl_propensity_leads"),
])
def test_lead_use_cases_route_to_separate_outputs(client, question, source_table):
    response = client.post("/ml/leads/query", headers=_headers(client, "bankuser0001"), json={"question": question})
    assert response.status_code == 200
    assert response.json()["source_table"] == source_table
    assert response.json()["total_count"] > 0
    record = response.json()["records"][0]
    assert record["business_unit"] in {"TIER_1", "TIER_2", "TIER_3", "TIER_4", "PLTB", "SALPL"}
    if source_table == "ml_pl_propensity_leads":
        assert record["preferred_contact_channel"] in {"CALL", "SMS", "EMAIL", "WHATSAPP"}
        assert record["preferred_contact_time"]
        assert record["AA_LAST_6M_AVG_BALANCE"] > 0
        assert record["AA_LAST_6M_DEBIT_AMOUNT"] > 0
        assert record["AA_LAST_6M_CREDIT_AMOUNT"] > 0
        assert record["AA_BASED_OFFER_AMOUNT"] > 0
        assert record["offer_amount"] > 0
        assert 300 <= record["bureau_cibil_score"] <= 900
        assert record["bureau_enquiries_6m"] >= 0
        assert record["bureau_active_external_loans"] >= 0
        assert record["bureau_current_exposure"] >= 0
        assert record["bureau_bounces_6m"] >= 0
        assert record["bureau_max_dpd_6m"] >= 0
    if source_table == "ml_pl_risk_review_leads":
        assert record["portfolio_segment_risk_probability"] > 0
        assert record["risk_segment"] in {"Super Red", "Red", "Risk Review"}
        assert record["review_status"].startswith("Risk probability ")
        assert "risk_team_tag" not in record
        assert "reason_code" not in record


@pytest.mark.integration
def test_risk_scorecard_segments_are_stored_and_exported(client):
    headers = _headers(client, "bankuser0001")
    question = "Who are the risky customers in Mumbai branch?"
    response = client.post("/ml/leads/query", headers=headers, json={"question": question})
    assert response.status_code == 200
    records = response.json()["records"]
    assert {"Red", "Super Red"}.issubset({row["risk_segment"] for row in records})
    assert all("risk_team_tag" not in row and "reason_code" not in row for row in records)

    export = client.post("/ml/leads/export", headers=headers, json={"question": question})
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content), read_only=True)
    columns = next(workbook.active.iter_rows(values_only=True))
    assert {"risk_segment", "portfolio_segment_risk_probability", "review_status"}.issubset(columns)
    assert "risk_team_tag" not in columns
    assert "reason_code" not in columns


@pytest.mark.integration
def test_propensity_ids_names_and_delinquency_are_consistent(client):
    headers = _headers(client, "bankuser0001")
    campaign = client.post("/ml/leads/query", headers=headers, json={
        "question": "Show hot personal loan leads for Mumbai branch",
    }).json()["records"]
    propensity = client.post("/ml/leads/query", headers=headers, json={
        "question": "Show PL propensity bands in Mumbai branch",
    }).json()["records"]
    campaign_names = {row["customer_ref"]: row["customer_name"] for row in campaign}
    assert all(campaign_names[row["customer_ref"]] == row["customer_name"] for row in propensity)
    delayed = [row for row in propensity if row["bureau_max_dpd_6m"] == 30]
    assert delayed and all(row["bureau_bounces_6m"] >= 1 for row in delayed)
    assert not any(key.startswith("aa_") for row in propensity for key in row)
    positive_net = [row for row in propensity if row["AA_LAST_6M_CREDIT_AMOUNT"] > row["AA_LAST_6M_DEBIT_AMOUNT"]]
    non_positive_net = [row for row in propensity if row["AA_LAST_6M_CREDIT_AMOUNT"] <= row["AA_LAST_6M_DEBIT_AMOUNT"]]
    assert positive_net and non_positive_net
    assert all(row["AA_BASED_OFFER_AMOUNT"] > row["offer_amount"] for row in positive_net)
    assert all(row["AA_BASED_OFFER_AMOUNT"] == row["offer_amount"] for row in non_positive_net)

    export = client.post("/ml/leads/export", headers=headers, json={
        "question": "Show PL propensity bands in Mumbai branch",
    })
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content), read_only=True)
    headers = next(workbook.active.iter_rows(values_only=True))
    assert {"bureau_cibil_score", "offer_amount", "AA_BASED_OFFER_AMOUNT",
            "AA_LAST_6M_AVG_BALANCE", "AA_LAST_6M_DEBIT_AMOUNT",
            "AA_LAST_6M_CREDIT_AMOUNT"}.issubset(headers)
    assert "average_ticket_size" not in headers


@pytest.mark.integration
def test_all_propensity_bands_and_explicit_band_queries(client):
    headers = _headers(client, "bankuser0001")
    all_bands = client.post("/ml/leads/query", headers=headers, json={
        "question": "What are the distinct propensity bands for the bank as a whole?",
    })
    assert all_bands.status_code == 200
    result = all_bands.json()
    assert result["total_count"] == 3_072
    assert result["band_counts"] == {
        "VERY_HIGH": 512, "HIGH": 1_024, "MEDIUM": 1_024, "LOW": 512,
    }
    for phrase, band, expected in (
        ("very high", "VERY_HIGH", 512),
        ("high", "HIGH", 1_024),
        ("medium", "MEDIUM", 1_024),
        ("low", "LOW", 512),
    ):
        response = client.post("/ml/leads/query", headers=headers, json={
            "question": f"Show {phrase} propensity customers for the bank as a whole",
        })
        assert response.status_code == 200
        assert response.json()["total_count"] == expected
        assert all(row["propensity_band"] == band for row in response.json()["records"])


@pytest.mark.security
def test_branch_cannot_read_mumbai_or_bank_wide_leads(client):
    headers = _headers(client, "bankuser0137")
    for question in (
        "Show hot PL leads for Mumbai branch",
        "Show hot PL leads for the bank as a whole",
    ):
        response = client.post("/ml/leads/query", headers=headers, json={"question": question})
        assert response.status_code == 403
        export = client.post("/ml/leads/export", headers=headers, json={"question": question})
        assert export.status_code == 403
